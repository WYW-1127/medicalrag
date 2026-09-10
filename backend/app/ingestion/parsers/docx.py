from pathlib import Path

from docx import Document

from app.ingestion.models import ParsedDocument
from app.ingestion.registry import file_hash
from app.ingestion.tree import SectionTreeBuilder


class DocxParser:
    """python-docx：style 名 Heading N → 层级 N；内嵌表格逐个转 markdown 风格文本。"""

    def parse(self, path: Path, *, department: str, source_path: str) -> ParsedDocument:
        builder = SectionTreeBuilder()
        d = Document(str(path))
        for p in d.paragraphs:
            style_name = p.style.name if p.style is not None else ""
            style = style_name.lower()
            text = p.text.strip()
            if not text:
                continue
            if style.startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:
                    level = 2
                builder.add_heading(level, text)
            else:
                builder.add_text(text)
        for i, tbl in enumerate(d.tables, 1):
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in tbl.rows]
            builder.add_table("\n".join(rows), f"表格(#{i})")
        return ParsedDocument(
            source_path=source_path,
            doc_type="guideline",
            department=department,
            title=path.stem,
            doc_hash=file_hash(path),
            sections=builder.build(),
        )

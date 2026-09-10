import json
from pathlib import Path

from docx import Document

from app.ingestion.parsers import get_parser
from app.ingestion.parsers.docx import DocxParser
from app.ingestion.parsers.html import HtmlParser
from app.ingestion.parsers.structured import StructuredParser


def test_html_parser_extracts_markdown(tmp_path: Path):
    html = """<html><body><article>
    <h1>冠心病诊疗指南</h1><p>冠心病的主要危险因素包括高血压与糖尿病。</p>
    <h2>诊断</h2><p>依据典型心绞痛症状与心电图改变。</p>
    </article></body></html>"""
    f = tmp_path / "g.html"
    f.write_text(html, encoding="utf-8")
    doc = HtmlParser().parse(f, department="心血管", source_path="x.html")
    assert doc.title == "冠心病诊疗指南"
    # h1 与 h2 之间的正文挂在 h1 节点自身
    assert "高血压" in doc.sections[0].text
    children = [c for c in doc.sections[0].children if not c.is_table]
    assert children[0].title == "诊断"
    assert "心绞痛" in children[0].text


def test_docx_parser_heading_and_table(tmp_path: Path):
    d = Document()
    d.add_heading("糖尿病指南", 1)
    d.add_heading("血糖控制目标", 2)
    d.add_paragraph("空腹血糖 4.4-7.0 mmol/L。")
    tbl = d.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "指标"
    tbl.cell(0, 1).text = "目标"
    tbl.cell(1, 0).text = "HbA1c"
    tbl.cell(1, 1).text = "<7%"
    f = tmp_path / "d.docx"
    d.save(str(f))
    doc = DocxParser().parse(f, department="内分泌", source_path="d.docx")

    def walk(nodes):
        for n in nodes:
            yield n
            yield from walk(n.children)

    assert doc.sections[0].title == "糖尿病指南"
    assert any(s.is_table and "HbA1c" in s.text for s in walk(doc.sections))


def test_json_drug_label(tmp_path: Path):
    data = {
        "name": "阿司匹林肠溶片",
        "department": "心血管",
        "fields": {"适应证": "心肌梗死、心绞痛", "禁忌": "活动性消化道出血"},
    }
    f = tmp_path / "asp.json"
    f.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    doc = StructuredParser().parse(f, department="综合", source_path="a.json")
    assert doc.doc_type == "drug_label"
    assert doc.title == "阿司匹林肠溶片"
    assert doc.department == "心血管"
    fields = {c.title: c.text for c in doc.sections[0].children}
    assert fields["禁忌"] == "活动性消化道出血"


def test_csv_records_grouped(tmp_path: Path):
    f = tmp_path / "icd.csv"
    f.write_text("编码,名称\nI10,原发性高血压\nE11,2型糖尿病\n", encoding="utf-8-sig")
    doc = StructuredParser().parse(f, department="综合", source_path="i.csv")
    assert doc.doc_type == "structured_table"
    assert "I10" in doc.sections[0].children[0].text


def test_get_parser_dispatch(tmp_path: Path):
    assert get_parser(tmp_path / "a.md").__class__.__name__ == "MarkdownParser"
    assert get_parser(tmp_path / "a.pdf").__class__.__name__ == "PdfParser"
    assert get_parser(tmp_path / "a.json").__class__.__name__ == "StructuredParser"
    try:
        get_parser(tmp_path / "a.txt")
        raise AssertionError("should raise ValueError")
    except ValueError:
        pass
